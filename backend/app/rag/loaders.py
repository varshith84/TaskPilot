"""
Document loaders for supported file formats.
Normalizes all inputs to LangChain Document objects.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Iterator

from langchain_core.documents import Document

def load_document(file_path: str, file_type: str, document_id: str, original_filename: str) -> list[Document]:
    """
    Load a file and return LangChain Documents with unified metadata.
    """
    if file_type == "pdf":
        return _load_pdf(file_path, document_id, original_filename)
    elif file_type == "docx":
        return _load_docx(file_path, document_id, original_filename)
    elif file_type in ("txt", "md"):
        return _load_text(file_path, file_type, document_id, original_filename)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

def _load_pdf(file_path: str, document_id: str, filename: str) -> list[Document]:
    from langchain_community.document_loaders import PyPDFLoader
    
    loader = PyPDFLoader(file_path)
    docs = []
    for i, page in enumerate(loader.lazy_load()):
        # page.metadata typically has 'page' (0-indexed or similar depending on version).
        # We enforce a clean 1-indexed page.
        page_num = i + 1
        docs.append(Document(
            page_content=page.page_content,
            metadata={
                "document_id": document_id,
                "filename": filename,
                "file_type": "pdf",
                "page": page_num
            }
        ))
    return docs

def _load_docx(
    file_path: str,
    document_id: str,
    filename: str
) -> list[Document]:
    from docx import Document as DocxDocument

    docx_file = DocxDocument(file_path)

    text_parts = []

    # Extract normal paragraphs
    for paragraph in docx_file.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # Extract text from tables
    for table in docx_file.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    text_parts.append(text)

    content = "\n".join(text_parts).strip()

    if not content:
        raise ValueError(
            "No extractable text found in DOCX. "
            "The document may contain only images, scans, or unsupported text boxes."
        )

    return [
        Document(
            page_content=content,
            metadata={
                "document_id": document_id,
                "filename": filename,
                "file_type": "docx",
            },
        )
    ]

def _load_text(file_path: str, file_type: str, document_id: str, filename: str) -> list[Document]:
    # Read text safely handling potential encoding issues
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            content = f.read()
            
    return [Document(
        page_content=content,
        metadata={
            "document_id": document_id,
            "filename": filename,
            "file_type": file_type
        }
    )]
